"""
Export routes - Download cleaned data as CSV/Excel, generate PDF report
"""

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel
from typing import Optional
import io, csv, json
from routes.upload import get_session_data

router = APIRouter()


@router.get("/{session_id}/csv")
async def export_csv(session_id: str):
    """Download cleaned dataset as CSV"""
    s = get_session_data(session_id)
    headers = s["headers"]
    rows = s["rows"]

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow({h: row.get(h, "") for h in headers})

    output.seek(0)
    filename = s["filename"].rsplit(".", 1)[0] + "_cleaned.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/{session_id}/excel")
async def export_excel(session_id: str):
    """Download cleaned dataset as Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl not installed. Run: pip install openpyxl")

    s = get_session_data(session_id)
    headers = s["headers"]
    rows = s["rows"]
    an = s["analysis"]

    wb = openpyxl.Workbook()

    # ── Sheet 1: Cleaned Data ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Cleaned Data"

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    alt_fill = PatternFill(start_color="F0F4FF", end_color="F0F4FF", fill_type="solid")

    for ci, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(headers, 1):
            cell = ws1.cell(row=ri, column=ci, value=row.get(h, ""))
            if ri % 2 == 0:
                cell.fill = alt_fill

    # Auto-width columns
    for ci, h in enumerate(headers, 1):
        max_len = max(len(str(h)), max((len(str(row.get(h, ""))) for row in rows[:100]), default=0))
        ws1.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 40)

    # ── Sheet 2: Statistics ─────────────────────────────────────────
    ws2 = wb.create_sheet("Statistics")
    stats = an.get("stats", {})
    num_cols = an.get("num_cols", [])

    stat_headers = ["Column", "Count", "Missing", "Mean", "Median", "Std Dev", "Min", "Max", "Outlier %", "Skewness"]
    for ci, h in enumerate(stat_headers, 1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for ri, col in enumerate(num_cols, 2):
        st = stats.get(col, {})
        vals = [col, st.get("count"), st.get("missing"), st.get("mean"), st.get("median"),
                st.get("std"), st.get("min"), st.get("max"), st.get("outlier_pct"), st.get("skewness")]
        for ci, v in enumerate(vals, 1):
            ws2.cell(row=ri, column=ci, value=v)

    for ci in range(1, len(stat_headers) + 1):
        ws2.column_dimensions[get_column_letter(ci)].width = 14

    # ── Sheet 3: Summary ───────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")
    brand = s.get("brand", {})
    summary_data = [
        ["DataMind Enterprise — Analysis Report"],
        [],
        ["File", s["filename"]],
        ["Dataset Type", brand.get("name", "General")],
        ["Total Records", s["row_count"]],
        ["Original Records", s["original_row_count"]],
        ["Columns", len(headers)],
        ["Numeric Columns", len(num_cols)],
        ["Categorical Columns", len(an.get("cat_cols", []))],
        ["Data Quality Score", f"{an.get('quality_score', 0)}/100"],
        ["Duplicates Removed", an.get("dupes", 0)],
        ["Missing Values Filled", an.get("missing_cells", 0)],
        ["Outliers Detected", an.get("outlier_cells", 0)],
    ]
    for ri, row_data in enumerate(summary_data, 1):
        for ci, v in enumerate(row_data, 1):
            cell = ws3.cell(row=ri, column=ci, value=v)
            if ri == 1:
                cell.font = Font(bold=True, size=14)
            elif len(row_data) == 2 and ci == 1:
                cell.font = Font(bold=True)

    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 30

    # Save
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = s["filename"].rsplit(".", 1)[0] + "_analyzed.xlsx"
    return StreamingResponse(
        iter([output.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


class ReportExportRequest(BaseModel):
    session_id: str
    report: dict  # The generated report sections
    user_name: Optional[str] = "Analyst"
    company: Optional[str] = "DataMind"


@router.post("/report/html")
async def export_report_html(req: ReportExportRequest):
    """Generate a printable HTML report (open in browser → Print → Save as PDF)"""
    s = get_session_data(req.session_id)
    an = s["analysis"]
    brand = s.get("brand", {})
    now_str = __import__("datetime").datetime.now().strftime("%B %d, %Y")

    def md_to_html(text: str) -> str:
        if not text:
            return "<p>Not generated.</p>"
        import re
        text = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', text)
        text = re.sub(r'\*(.+?)\*', r'<em>\1</em>', text)
        text = re.sub(r'^### (.+)$', r'<h4>\1</h4>', text, flags=re.MULTILINE)
        text = re.sub(r'^## (.+)$', r'<h3>\1</h3>', text, flags=re.MULTILINE)
        text = re.sub(r'^# (.+)$', r'<h2>\1</h2>', text, flags=re.MULTILINE)
        text = re.sub(r'^🔴 (.+)$', r'<div class="risk high">🔴 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^🟡 (.+)$', r'<div class="risk med">🟡 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^🟢 (.+)$', r'<div class="risk low">🟢 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^📈 (.+)$', r'<div class="scenario opt">📈 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^📊 (.+)$', r'<div class="scenario base">📊 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^📉 (.+)$', r'<div class="scenario pess">📉 \1</div>', text, flags=re.MULTILINE)
        text = re.sub(r'^• (.+)$', r'<li>\1</li>', text, flags=re.MULTILINE)
        text = re.sub(r'^- (.+)$', r'<li>\1</li>', text, flags=re.MULTILINE)
        text = re.sub(r'^(\d+)\. (.+)$', r'<li><strong>\1.</strong> \2</li>', text, flags=re.MULTILINE)
        text = re.sub(r'(<li>.*?</li>\n?)+', lambda m: f'<ul>{m.group(0)}</ul>', text, flags=re.DOTALL)
        text = text.replace('\n\n', '</p><p>').replace('\n', '<br>')
        return f"<p>{text}</p>"

    sections = [
        ("01", "Executive Summary", "executive"),
        ("02", "KPI Analysis", "kpis"),
        ("03", "Trend & Pattern Analysis", "trends"),
        ("04", "Risk Assessment", "risks"),
        ("05", "Strategic Recommendations", "recommendations"),
        ("06", "Forecast & Outlook", "forecast"),
    ]

    sections_html = ""
    for num, title, key in sections:
        content = md_to_html(req.report.get(key, ""))
        sections_html += f"""
        <div class="section">
            <div class="sec-header">
                <span class="sec-num">{num}</span>
                <h2>{title}</h2>
            </div>
            <div class="sec-body">{content}</div>
        </div>"""

    stats_rows = ""
    stats = an.get("stats", {})
    for col in an.get("num_cols", [])[:15]:
        st = stats.get(col, {})
        stats_rows += f"""<tr>
            <td><strong>{col}</strong></td>
            <td>{st.get('count','')}</td><td>{st.get('missing','')}</td>
            <td>{st.get('mean','')}</td><td>{st.get('median','')}</td>
            <td>{st.get('std','')}</td><td>{st.get('min','')}</td>
            <td>{st.get('max','')}</td><td>{st.get('outlier_pct','')}%</td>
        </tr>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>DataMind Report — {s['filename']}</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: 'Georgia', serif; color: #0a1020; background: #fff; font-size: 14px; line-height: 1.8; }}
  .cover {{ min-height: 100vh; background: linear-gradient(135deg, #07090F 0%, #0F1A2E 100%); color: #fff; padding: 60px 72px; display: flex; flex-direction: column; justify-content: space-between; page-break-after: always; }}
  .cover-eyebrow {{ font-size: 11px; letter-spacing: 4px; text-transform: uppercase; color: #3B82F6; margin-bottom: 24px; font-weight: 700; }}
  .cover-title {{ font-size: 48px; font-weight: 800; line-height: 1.1; margin-bottom: 16px; }}
  .cover-meta {{ font-size: 16px; color: #8BA0B8; margin-bottom: 40px; }}
  .cover-footer {{ display: flex; justify-content: space-between; padding-top: 24px; border-top: 1px solid #1E2D3D; font-size: 13px; color: #8BA0B8; }}
  .quality-bar {{ margin: 36px 72px; padding: 28px 32px; border: 2px solid #3B82F6; border-radius: 14px; display: flex; align-items: center; gap: 40px; }}
  .q-score {{ font-size: 56px; font-weight: 900; color: {'#10B981' if an.get('quality_score',0) > 80 else '#F59E0B'}; line-height: 1; }}
  .q-label {{ font-size: 12px; color: #666; }}
  .q-details {{ display: flex; gap: 32px; }}
  .q-detail {{ text-align: center; }}
  .q-detail-val {{ font-size: 24px; font-weight: 800; color: #0a1020; }}
  .q-detail-label {{ font-size: 11px; color: #666; }}
  .section {{ margin: 0 72px; padding: 32px 0; border-bottom: 1px solid #e0e8f0; }}
  .sec-header {{ display: flex; align-items: baseline; gap: 16px; margin-bottom: 20px; }}
  .sec-num {{ font-size: 10px; letter-spacing: 3px; color: #3B82F6; font-weight: 700; font-family: monospace; min-width: 28px; }}
  h2 {{ font-size: 24px; font-weight: 700; color: #0a1020; }}
  h3 {{ font-size: 18px; font-weight: 700; color: #0A3054; margin: 16px 0 8px; }}
  h4 {{ font-size: 15px; font-weight: 700; color: #0A3054; margin: 14px 0 6px; }}
  .sec-body {{ padding-left: 44px; }}
  .sec-body p {{ margin: 8px 0; }}
  .sec-body ul {{ padding-left: 20px; margin: 8px 0; list-style: disc; }}
  .sec-body li {{ margin: 6px 0; }}
  .risk {{ padding: 8px 14px; border-radius: 8px; margin: 6px 0; font-weight: 500; }}
  .risk.high {{ background: #FEE2E2; color: #991B1B; }}
  .risk.med  {{ background: #FEF3C7; color: #92400E; }}
  .risk.low  {{ background: #D1FAE5; color: #065F46; }}
  .scenario {{ padding: 8px 14px; border-radius: 8px; margin: 6px 0; }}
  .scenario.opt  {{ background: #D1FAE5; color: #065F46; }}
  .scenario.base {{ background: #DBEAFE; color: #1E40AF; }}
  .scenario.pess {{ background: #FEE2E2; color: #991B1B; }}
  .appendix {{ margin: 0 72px; padding: 32px 0; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 12px; margin-top: 12px; }}
  th {{ background: #1E3A5F; color: #fff; padding: 9px 12px; text-align: left; font-weight: 700; }}
  td {{ padding: 8px 12px; border-bottom: 1px solid #DDE6F5; }}
  tr:nth-child(even) td {{ background: #F0F4FF; }}
  .footer {{ margin: 32px 72px; padding-top: 20px; border-top: 1px solid #eee; font-size: 11px; color: #999; display: flex; justify-content: space-between; }}
  @media print {{ body {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }} .cover {{ page-break-after: always; }} }}
</style>
</head>
<body>

<div class="cover">
  <div>
    <div class="cover-eyebrow">◈ DataMind Enterprise · Intelligence Report</div>
    <div class="cover-title">{s['filename'].rsplit('.',1)[0]}</div>
    <div class="cover-meta">📊 {brand.get('name','General')} &nbsp;·&nbsp; {s['row_count']:,} Records &nbsp;·&nbsp; {len(s['headers'])} Dimensions &nbsp;·&nbsp; Quality {an.get('quality_score',0)}/100</div>
  </div>
  <div class="cover-footer">
    <div>Prepared for: {req.user_name} · {req.company}</div>
    <div>{now_str} &nbsp;·&nbsp; CONFIDENTIAL</div>
  </div>
</div>

<div class="quality-bar">
  <div>
    <div class="q-score">{an.get('quality_score',0)}</div>
    <div class="q-label">Quality Score /100</div>
  </div>
  <div class="q-details">
    <div class="q-detail"><div class="q-detail-val">{s['row_count']:,}</div><div class="q-detail-label">Records Analyzed</div></div>
    <div class="q-detail"><div class="q-detail-val">{an.get('dupes',0)}</div><div class="q-detail-label">Duplicates Removed</div></div>
    <div class="q-detail"><div class="q-detail-val">{an.get('missing_cells',0)}</div><div class="q-detail-label">Missing Values Filled</div></div>
    <div class="q-detail"><div class="q-detail-val">{an.get('outlier_cells',0)}</div><div class="q-detail-label">Outliers Detected</div></div>
  </div>
</div>

{sections_html}

<div class="appendix">
  <div class="sec-header"><span class="sec-num">07</span><h2>Statistical Appendix</h2></div>
  <table>
    <thead><tr><th>Column</th><th>Count</th><th>Missing</th><th>Mean</th><th>Median</th><th>Std Dev</th><th>Min</th><th>Max</th><th>Outlier %</th></tr></thead>
    <tbody>{stats_rows}</tbody>
  </table>
</div>

<div class="footer">
  <span>DataMind Enterprise · AI-Powered by Ollama · Generated {now_str}</span>
  <span>{req.user_name} · {req.company} · CONFIDENTIAL</span>
</div>

<script>window.addEventListener('load', () => window.print());</script>
</body>
</html>"""

    return Response(content=html, media_type="text/html")
